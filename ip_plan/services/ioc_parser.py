"""Извлечение депрофанированных индикаторов компрометации (IOC) из писем/
бюллетеней регулятора (PDF/DOCX) — см. docs/plans/2026-09-07-ioc-parser-design.md.

Только текстовый разбор, никаких сетевых обращений и никакого "оживления"
найденных адресов в кликабельные ссылки — сами письма прямо запрещают
переход по этим ссылкам, парсер не должен даже теоретически предоставлять
такую возможность (см. refang() ниже: чисто строковая замена в текстовом
поле результата).

Framework-agnostic: не импортирует Django, пригоден и для management-команды
(ioc_scan), и для вызова из view (regulator_entry_ioc_parse) напрямую на
файлах из request.FILES, без сохранения на диск.
"""

import logging
import re
from dataclasses import dataclass, field

logger = logging.getLogger(__name__)

# ── Регулярные выражения (см. design doc — проверены на реальных письмах) ──

URL_RE = re.compile(r"hxxps?\[:\]//[A-Za-z0-9\[\].\-_/%:]+", re.IGNORECASE)
IP_RE = re.compile(r"\b(?:\d{1,3}\[\.\]){3}\d{1,3}\b")
HASH_RE = re.compile(r"\b[a-fA-F0-9]{64}\b")
DOMAIN_RE = re.compile(
    r"\b[a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?"
    r"(?:\[\.\][a-zA-Z0-9](?:[a-zA-Z0-9-]{0,61}[a-zA-Z0-9])?)+\b"
)
# Хвостовая пунктуация ПРЕДЛОЖЕНИЯ, приклеившаяся к найденному токену
# ("...msi." -> "...msi") — не путать с [.]/[:] внутри самого депрофанированного
# адреса, этот regex режет только с САМОГО КОНЦА строки один раз.
TRAIL_PUNCT_RE = re.compile(r"[.;:,)\]]+$")
# "Похоже на IP" (все группы — числа) — чтобы отфильтровать из результатов
# DOMAIN_RE случайно совпавшие числовые последовательности (см. design doc,
# шаг 4: домен из одних цифр — это на самом деле IP, не домен).
IP_SHAPED_RE = re.compile(r"^\d+(\[\.\]\d+)+$")
# "Похоже на IP с приклеенной сноской" — 1-2 латинские буквы сразу после
# последнего октета без разделителя, например "89[.]221[.]203[.]162l" (перед
# ";"). Обычный IP_RE такое НЕ ловит вообще: \b после последнего \d{1,3} не
# матчится, если следующий символ — буква (тоже \w, между двумя \w нет
# границы слова) — отсюда отдельный, более широкий паттерн специально под
# этот сценарий, см. design doc.
IP_FOOTNOTE_RE = re.compile(r"\b(?:\d{1,3}\[\.\]){3}\d{1,3}[A-Za-z]{1,2}\b")

# Символы, которыми ЗАКАНЧИВАЕТСЯ строка, если это последняя строка
# самостоятельного пункта списка (не разрыв URL посередине) — см.
# join_broken_url_lines.
LINE_TERMINAL_CHARS = ";.,)"

SUPPORTED_SUFFIXES = (".pdf", ".docx")


# ── Дата-класс индикатора ────────────────────────────────────────────────

@dataclass
class Indicator:
    kind: str  # "ip" | "domain" | "url" | "sha256"
    value: str  # рефанированное (нормальное) значение
    value_defanged: str | None  # None для sha256 — там нет депрофанации
    source_files: set = field(default_factory=set)
    occurrences: int = 0
    needs_review: bool = False
    note: str = ""

    def to_dict(self):
        d = {
            "type": self.kind,
            "value": self.value,
            "source_files": sorted(self.source_files),
            "occurrences": self.occurrences,
            "needs_review": self.needs_review,
        }
        if self.value_defanged is not None:
            d["value_defanged"] = self.value_defanged
        if self.note:
            d["note"] = self.note
        return d


def refang(defanged: str) -> str:
    """Только текстовая замена — НЕ создаёт кликабельную ссылку, ничего не
    резолвит и никуда не обращается по сети (см. модульный docstring)."""
    return (
        defanged.replace("[.]", ".").replace("[:]", ":")
        .replace("hxxps", "https").replace("hxxp", "http")
    )


def _validate_ip_octets(ip_plain: str) -> bool:
    parts = ip_plain.split(".")
    if len(parts) != 4:
        return False
    try:
        return all(0 <= int(p) <= 255 for p in parts)
    except ValueError:
        return False


# ── Склейка строк, разорванных переносом посередине URL ─────────────────

def join_broken_url_lines(lines):
    """lines — список строк ОДНОГО файла в исходном порядке (страницы PDF
    подряд, либо параграфы/ячейки таблиц DOCX). Возвращает новый список,
    где строки, разорванные переносом посреди hxxp-URL, склеены в одну (см.
    design doc про правило: только для строк с 'hxxp', без терминальной
    пунктуации на конце, если следующая строка начинается с латинской
    буквы/цифры — приклеить первое "слово" следующей строки без пробела,
    остаток — как отдельную строку дальше).

    НЕ трогает обычные построчные списки (IP/домены/хэши без ';' на конце)
    — они не содержат 'hxxp', условие склейки на них не срабатывает вообще."""
    work = list(lines)  # рабочая копия — мутируем на месте (заменяем/удаляем)
    result = []
    i = 0
    while i < len(work):
        current = work[i]
        while True:
            stripped = current.rstrip()
            if not stripped or "hxxp" not in stripped.lower():
                break
            if stripped[-1] in LINE_TERMINAL_CHARS:
                break
            if i + 1 >= len(work):
                break
            next_lstripped = work[i + 1].lstrip()
            if not next_lstripped:
                break
            first_char = next_lstripped[0]
            if not (first_char.isascii() and first_char.isalnum()):
                break
            parts = next_lstripped.split(None, 1)
            first_word = parts[0]
            remainder = parts[1] if len(parts) > 1 else ""
            current = stripped + first_word
            if remainder:
                work[i + 1] = remainder
            else:
                del work[i + 1]
            # Продолжаем — URL мог перенестись больше одного раза подряд.
        result.append(current)
        i += 1
    return result


# ── Извлечение индикаторов из текста одного файла ────────────────────────

def extract_indicators_from_text(lines, source_file, indicators):
    """lines — уже склеенный (join_broken_url_lines) список строк файла.
    indicators — dict {kind: {match_key: Indicator}}, мутируется на месте
    (глобальная дедупликация по всей папке — вызывающий код передаёт один и
    тот же dict для всех файлов прогона)."""
    full_text = "\n".join(lines)

    def record(kind, value_plain, value_defanged, needs_review=False, note=""):
        bucket = indicators.setdefault(kind, {})
        key = value_plain
        ind = bucket.get(key)
        if ind is None:
            ind = Indicator(kind=kind, value=value_plain, value_defanged=value_defanged)
            bucket[key] = ind
        ind.source_files.add(source_file)
        ind.occurrences += 1
        if needs_review:
            ind.needs_review = True
            if note and note not in ind.note:
                ind.note = f"{ind.note}; {note}" if ind.note else note

    # 1. SHA256 — на исходном (не депрофанированном) тексте, независимо от
    #    остального пайплайна (хэши не пересекаются с URL/IP/domain regex).
    for m in HASH_RE.finditer(full_text):
        h = m.group(0).lower()
        record("sha256", h, None)

    # 2. URL — до IP/доменов, чтобы адрес внутри URL не задвоился отдельным
    #    индикатором (см. design doc, порядок применения регулярок).
    working = full_text
    for m in URL_RE.finditer(full_text):
        raw = TRAIL_PUNCT_RE.sub("", m.group(0))
        if not raw:
            continue
        record("url", refang(raw), raw)
    working = URL_RE.sub(" ", working)

    # 3. IP с "приклеенной сноской" — отдельным проходом ПЕРЕД строгим
    #    IP_RE (строгий паттерн такие вообще не матчит, см. IP_FOOTNOTE_RE).
    for m in IP_FOOTNOTE_RE.finditer(working):
        raw_full = m.group(0)
        trimmed = re.sub(r"[A-Za-z]{1,2}$", "", raw_full)
        plain = refang(trimmed)
        note = "possible footnote artifact trimmed"
        if not _validate_ip_octets(plain):
            note += "; invalid octet range after trim"
        record("ip", plain, trimmed, needs_review=True, note=note)
    working = IP_FOOTNOTE_RE.sub(" ", working)

    # 4. Обычные, чистые депрофанированные IP.
    for m in IP_RE.finditer(working):
        raw = m.group(0)
        plain = refang(raw)
        if _validate_ip_octets(plain):
            record("ip", plain, raw)
        else:
            record("ip", plain, raw, needs_review=True, note="invalid octet range (0-255)")
    working = IP_RE.sub(" ", working)

    # 5. Домены — из остатка, отфильтровав то, что на самом деле IP-подобно.
    for m in DOMAIN_RE.finditer(working):
        raw = TRAIL_PUNCT_RE.sub("", m.group(0))
        if not raw or IP_SHAPED_RE.match(raw):
            continue
        record("domain", refang(raw), raw)


# ── Извлечение текста из файлов ──────────────────────────────────────────

def extract_pdf_lines(file_obj_or_path):
    """Возвращает список строк (постранично, порядок сохранён) или None,
    если из PDF не удалось извлечь текст вообще (защищён паролем, скан без
    текстового слоя и т.п.) — вызывающий код превращает None в warning и
    пропускает файл, не падая."""
    import pdfplumber

    lines = []
    try:
        with pdfplumber.open(file_obj_or_path) as pdf:
            any_text = False
            for page in pdf.pages:
                text = page.extract_text() or ""
                if text:
                    any_text = True
                lines.extend(text.split("\n"))
            if not any_text:
                return None
    except Exception as exc:  # pdfminer/pdfplumber кидают разные типы на битые/защищённые PDF
        logger.warning("Не удалось прочитать PDF: %s", exc)
        return None
    return lines


def extract_docx_lines(file_obj_or_path):
    """document.paragraphs + document.tables (в этом порядке — сначала весь
    основной текст, потом содержимое таблиц), см. design doc. None — файл
    повреждён/не открылся как DOCX."""
    import docx

    try:
        document = docx.Document(file_obj_or_path)
    except Exception as exc:
        logger.warning("Не удалось прочитать DOCX: %s", exc)
        return None

    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    lines.append(p.text)
    return lines


def extract_lines(filename, file_obj_or_path):
    """Диспетчер по расширению имени файла. Возвращает (lines, warning) —
    warning не None, если файл пропущен (неподдерживаемый формат, ошибка
    чтения) — lines в этом случае пуст."""
    lower = filename.lower()
    if lower.endswith(".pdf"):
        lines = extract_pdf_lines(file_obj_or_path)
        if lines is None:
            return [], f"{filename}: не удалось извлечь текст (защищён паролем или без текстового слоя)"
        return lines, None
    if lower.endswith(".docx"):
        lines = extract_docx_lines(file_obj_or_path)
        if lines is None:
            return [], f"{filename}: файл повреждён или не открылся как DOCX"
        return lines, None
    return [], f"{filename}: неподдерживаемый формат (ожидается .pdf/.docx)"


# ── Обработка одного файла (для веб-эндпоинта — файл уже открыт в памяти) ─

def parse_file(filename, file_obj, indicators):
    """file_obj — файлоподобный объект (Django UploadedFile или открытый
    файл). Мутирует indicators (см. extract_indicators_from_text). Возвращает
    warning-строку или None."""
    lines, warning = extract_lines(filename, file_obj)
    if warning:
        return warning
    joined = join_broken_url_lines(lines)
    extract_indicators_from_text(joined, filename, indicators)
    return None


# ── CLI/management-командный вход: сканирование папки на диске ──────────

def scan_folder(folder, recursive=False, per_file=False):
    """Возвращает dict в формате итогового JSON (см. design doc). folder —
    Path.

    per_file — если True, каждый файл СНАЧАЛА разбирается в свой
    изолированный dict (свои Indicator-объекты — occurrences/source_files
    только по этому файлу, для вывода под ключом "per_file"), и уже эти же
    объекты сразу переиспользуются как вклад в глобальный dict (через
    _merge_indicators) — один проход разбора на файл в обоих режимах,
    разница только в том, сохраняем ли промежуточный per-file срез."""
    import datetime
    from pathlib import Path

    folder = Path(folder)
    pattern = "**/*" if recursive else "*"
    all_files = sorted(p for p in folder.glob(pattern) if p.is_file())

    indicators = {}
    warnings = []
    per_file_results = {} if per_file else None

    for path in all_files:
        if path.suffix.lower() not in SUPPORTED_SUFFIXES:
            warnings.append({"file": path.name, "message": "неподдерживаемый формат, пропущен"})
            logger.warning("%s: неподдерживаемый формат, пропущен", path.name)
            continue

        file_indicators = {}
        with open(path, "rb") as f:
            warning = parse_file(path.name, f, file_indicators)
        if warning:
            warnings.append({"file": path.name, "message": warning})
            continue

        if per_file and any(file_indicators.values()):
            per_file_results[path.name] = {
                kind: [ind.to_dict() for ind in bucket.values()]
                for kind, bucket in file_indicators.items() if bucket
            }
        _merge_indicators(indicators, file_indicators)

    files_with_indicators = set()
    for bucket in indicators.values():
        for ind in bucket.values():
            files_with_indicators |= ind.source_files

    result = {
        "generated_at": datetime.datetime.now(datetime.timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source_folder": str(folder),
        "files_scanned": len(all_files),
        "files_with_indicators": len(files_with_indicators),
        "indicators": {
            kind: [ind.to_dict() for ind in bucket.values()]
            for kind, bucket in indicators.items()
        },
        "warnings": warnings,
    }
    if per_file:
        result["per_file"] = per_file_results
    return result


def _merge_indicators(target, source):
    """Складывает indicators-dict source (обычно результат одного файла) в
    target (глобальный аккумулятор) — суммирует occurrences, объединяет
    source_files, needs_review "заражает" (если хоть где-то True — True)."""
    for kind, bucket in source.items():
        target_bucket = target.setdefault(kind, {})
        for key, ind in bucket.items():
            existing = target_bucket.get(key)
            if existing is None:
                target_bucket[key] = ind
                continue
            existing.source_files |= ind.source_files
            existing.occurrences += ind.occurrences
            existing.needs_review = existing.needs_review or ind.needs_review
            if ind.note and ind.note not in existing.note:
                existing.note = f"{existing.note}; {ind.note}" if existing.note else ind.note
